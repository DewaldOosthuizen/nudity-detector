"""Extended tests for src/reporting/report_manager.py to boost coverage."""
import json
import os

import pytest

from src.core.models import ReportEntry, ScanConfig, SessionState
from src.reporting.report_manager import ReportManager


def _make_entry(**kwargs):
    defaults = {
        "file": "/tmp/img.jpg",
        "media_type": "image",
        "model_name": "nudenet",
        "threshold_percent": 60.0,
        "confidence_percent": 0.0,
        "nudity_detected": False,
        "detected_classes": "[]",
        "thumbnail": "",
        "date_classified": "2024-01-01 00:00:00",
    }
    defaults.update(kwargs)
    return ReportEntry.from_dict(defaults)


# ---------------------------------------------------------------------------
# save_entries / load_entries round-trip
# ---------------------------------------------------------------------------

def test_save_and_load_entries(tmp_path):
    report = str(tmp_path / "report.xlsx")
    entries = [_make_entry(file=str(tmp_path / "a.jpg"), nudity_detected=True)]
    result = ReportManager.save_entries(entries, report)
    assert result is True
    assert os.path.exists(report)

    loaded = ReportManager.load_entries(report)
    assert len(loaded) == 1
    assert loaded[0].file == str(tmp_path / "a.jpg")


def test_save_entries_overwrites(tmp_path):
    """Saving twice should overwrite, not append."""
    report = str(tmp_path / "report.xlsx")
    entries1 = [_make_entry(file="/tmp/a.jpg")]
    entries2 = [_make_entry(file="/tmp/b.jpg"), _make_entry(file="/tmp/c.jpg")]
    ReportManager.save_entries(entries1, report)
    ReportManager.save_entries(entries2, report)
    loaded = ReportManager.load_entries(report)
    assert len(loaded) == 2


def test_load_entries_returns_empty_for_empty_xlsx(tmp_path):
    """An xlsx with only a header row and no data returns empty list."""
    import openpyxl
    report = str(tmp_path / "empty.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nudity Report"
    ws.append(["File"])
    wb.save(report)
    loaded = ReportManager.load_entries(report)
    assert loaded == []


# ---------------------------------------------------------------------------
# save_session / load_session round-trip
# ---------------------------------------------------------------------------

def test_save_and_load_session(tmp_path):
    report = str(tmp_path / "report.xlsx")
    session = SessionState(
        scan_config=ScanConfig(source_folder="/tmp"),
        results=[],
    )
    result = ReportManager.save_session(session, report)
    assert result is True

    loaded = ReportManager.load_session(report)
    assert isinstance(loaded, SessionState)
    assert loaded.scan_config.source_folder == "/tmp"


def test_load_session_missing_file():
    loaded = ReportManager.load_session("/nonexistent/report.xlsx")
    assert isinstance(loaded, SessionState)


def test_load_session_from_json(tmp_path):
    """Pass a .json path directly."""
    session = SessionState(scan_config=ScanConfig(source_folder="/data"))
    json_path = str(tmp_path / "sess.json")
    with open(json_path, "w") as f:
        json.dump(session.to_dict(), f)

    loaded = ReportManager.load_session(json_path)
    assert loaded.scan_config.source_folder == "/data"


def test_load_session_corrupt_json(tmp_path):
    """Corrupt JSON falls back to empty SessionState."""
    json_path = str(tmp_path / "corrupt_session.json")
    with open(json_path, "w") as f:
        f.write("{not valid json")

    loaded = ReportManager.load_session(json_path)
    assert isinstance(loaded, SessionState)


# ---------------------------------------------------------------------------
# create_demo_session
# ---------------------------------------------------------------------------

def test_create_demo_session_new_file(tmp_path):
    report = str(tmp_path / "report.xlsx")
    session = SessionState()
    result = ReportManager.create_demo_session(report, session)
    assert result is True
    assert os.path.exists(report)


def test_create_demo_session_existing_file(tmp_path):
    report = str(tmp_path / "report.xlsx")
    # First create
    ReportManager.save_entries([], report)
    session = SessionState(scan_config=ScanConfig(source_folder="/demo"))
    result = ReportManager.create_demo_session(report, session)
    assert result is True


# ---------------------------------------------------------------------------
# _embed_thumbnails — with actual base64 PNG thumbnail
# ---------------------------------------------------------------------------

def test_embed_thumbnails_with_valid_thumbnail(tmp_path):
    try:
        from PIL import Image as PILImage
    except ImportError:
        pytest.skip("PIL not available")

    # Generate a valid base64 thumbnail
    img = PILImage.new("RGB", (10, 10), color=(0, 0, 255))
    from io import BytesIO
    buf = BytesIO()
    img.save(buf, format="PNG")
    b64 = __import__("base64").b64encode(buf.getvalue()).decode()

    report = str(tmp_path / "report.xlsx")
    entry = _make_entry(thumbnail=b64)
    result = ReportManager.save_entries([entry], report)
    assert result is True


# ---------------------------------------------------------------------------
# validate_report_dir — system protected
# ---------------------------------------------------------------------------

def test_validate_report_dir_system_protected():
    valid, msg = ReportManager.validate_report_dir("/")
    assert valid is False
    assert "system directory" in msg.lower()


# ---------------------------------------------------------------------------
# formula injection / text-coercion regression tests
# ---------------------------------------------------------------------------

def test_save_entries_formula_leading_strings_are_text(tmp_path):
    """A value beginning with '=' must be stored as a literal text cell, not a formula."""
    import openpyxl

    report = str(tmp_path / "report.xlsx")
    entries = [
        _make_entry(
            file="=cmd|' /C calc'!Z9",
            media_type="=SUM(1,1)",
            model_name="+1",
            detected_classes='="x"',
        ),
    ]
    assert ReportManager.save_entries(entries, report) is True

    wb = openpyxl.load_workbook(report)
    ws = wb.active
    assert ws.cell(row=2, column=1).value == "=cmd|' /C calc'!Z9"
    assert ws.cell(row=2, column=1).data_type == "s"
    assert ws.cell(row=2, column=2).value == "=SUM(1,1)"
    assert ws.cell(row=2, column=2).data_type == "s"
    assert ws.cell(row=2, column=3).value == "+1"
    assert ws.cell(row=2, column=3).data_type == "s"
    assert ws.cell(row=2, column=7).value == '="x"'
    assert ws.cell(row=2, column=7).data_type == "s"


def test_save_and_load_entries_preserves_literal_formula_values(tmp_path):
    """Reloaded reports preserve formula-leading values as literal strings."""
    report = str(tmp_path / "report.xlsx")
    entries = [
        _make_entry(
            file="=cmd|' /C calc'!Z9",
            media_type="=SUM(1,1)",
            model_name="+1",
            detected_classes='="x"',
            nudity_detected=True,
        ),
    ]
    assert ReportManager.save_entries(entries, report) is True

    loaded = ReportManager.load_entries(report)
    assert len(loaded) == 1
    assert loaded[0].file == "=cmd|' /C calc'!Z9"
    assert loaded[0].media_type == "=SUM(1,1)"
    assert loaded[0].model_name == "+1"
    assert loaded[0].detected_classes == '="x"'


def test_saved_report_contains_no_formula_cells(tmp_path):
    """The saved xlsx must not evaluate any untrusted value as a formula."""
    import zipfile

    report = str(tmp_path / "report.xlsx")
    entries = [_make_entry(file="=cmd|' /C calc'!Z9", nudity_detected=True)]
    assert ReportManager.save_entries(entries, report) is True

    with zipfile.ZipFile(report) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")

    assert "<f>" not in xml, "a cell was written as a formula"
